#!/usr/bin/env python3
"""Importer: DMA-Excel (Skill-Output) -> Objektmodell.

Erzeugt schema-konforme Objekte aus ALLEN bewerteten IROs (wesentlich + nicht):
  - iro-*    je IRO (wesentlich ja/nein gem. 'Material?')
  - topic-*  je Thema (Standard + Sub-Thema); status = wesentlich, wenn >=1 IRO wesentlich
  - decision-dma-2026 (append-only), methodology-dma-2026, person-dma-lead

Nicht-wesentliche IROs/Themen werden bewusst MIT importiert — die Ausschluss-Begründung
ist prüfungsrelevant (Auditoren fechten gerade Ausschlüsse an).

Aufruf:  python3 tools/import_dma.py <pfad-zur-xlsx>
Benötigt: openpyxl.
"""
from __future__ import annotations
import sys, re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OBJ = ROOT / "objects"; OBJ.mkdir(parents=True, exist_ok=True)
STAND = "2026-06-09"
TYP = {"I-": "Impact-negativ", "I+": "Impact-positiv", "R": "Risk", "O": "Opportunity"}
UML = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss", "Ä": "ae", "Ö": "oe", "Ü": "ue"})


def slug(s, n=44):
    s = (s or "").translate(UML).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:n].strip("-")


def lvl(v):
    try:
        v = int(v)
    except (TypeError, ValueError):
        return "niedrig"
    return "hoch" if v >= 4 else "mittel" if v == 3 else "niedrig"


def num(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def q(s):
    return '"' + str(s or "").replace('"', "'").strip() + '"'


def write(fid, fm, body):
    lines = ["---"] + [f"{k}: {v}" for k, v in fm.items()] + ["---", "", body, ""]
    (OBJ / f"{fid}.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    if len(sys.argv) < 2:
        sys.exit("Aufruf: import_dma.py <xlsx>")
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    rows = [r for r in list(wb[wb.sheetnames[3]].iter_rows(values_only=True))[3:] if r and r[0]]

    write("person-dma-lead", {
        "id": "person-dma-lead", "type": "person", "owner": "person-dma-lead",
        "status": "aktiv", "stand": STAND, "vertraulichkeit": "intern", "rolle": "DMA Lead"},
        "# Person: DMA Lead\n\nPlatzhalter-Owner für die importierten DMA-2026-Objekte.")
    write("methodology-dma-2026", {
        "id": "methodology-dma-2026", "type": "methodology", "owner": "person-dma-lead",
        "status": "final", "stand": STAND, "review_zyklus": "P12M", "vertraulichkeit": "intern",
        "esrs_bezug": "ESRS-2-IRO-1"},
        "# Methodik: DMA 2026\n\nJe IRO: A (Scale/Magnitude), B (Scope/Probability), "
        "C (Irreversibility) -> Impact-Score; plus Financial-Score. Schwelle Score >= 3 = wesentlich.\n"
        "Mapping: Score 4-5 -> hoch, 3 -> mittel, 1-2 -> niedrig.")

    topics: dict[str, dict] = {}
    n_mat = 0
    for r in rows:
        std = str(r[1]).strip()
        sub = str(r[2]).strip()
        material = str(r[13]).strip().lower() == "ja"
        n_mat += material
        tid = f"topic-{slug(std)}-{slug(sub, 36)}"
        iid = f"iro-{slug(str(r[0]))}"
        t = topics.setdefault(tid, {"std": std, "sub": sub, "iros": [], "mat": False})
        t["iros"].append(iid)
        t["mat"] = t["mat"] or material
        write(iid, {
            "id": iid, "type": "iro", "owner": "person-dma-lead", "status": "final",
            "stand": STAND, "review_zyklus": "P12M", "vertraulichkeit": "intern",
            "esrs_bezug": std, "iro_typ": TYP.get(str(r[3]).strip(), str(r[3]).strip()),
            "impact_wesentlichkeit": lvl(r[11]), "finanz_wesentlichkeit": lvl(r[12]),
            "wesentlich": '"ja"' if material else '"nein"',
            "impact_score": num(r[11]), "financial_score": num(r[12]),
            "iro_nr": q(r[0]), "sub_thema": q(sub),
            "concerns": f"[{tid}]", "scored_under": "[methodology-dma-2026]"},
        f"# IRO {r[0]}: {sub}\n\n"
        f"**Typ:** {r[3]} · **Aktuell/Potenziell:** {r[4]} · **Wertschöpfungskette:** {r[5]} · "
        f"**Zeithorizont:** {r[6]} · **Wesentlich:** {'Ja' if material else 'Nein'}\n\n"
        f"## Beschreibung\n{r[7] or '—'}\n\n## Begründung\n{r[14] or '—'}")

    n_mat_topics = sum(1 for t in topics.values() if t["mat"])
    for tid, t in topics.items():
        write(tid, {
            "id": tid, "type": "topic", "owner": "person-dma-lead",
            "status": "wesentlich" if t["mat"] else "nicht-wesentlich",
            "stand": STAND, "review_zyklus": "P12M", "vertraulichkeit": "intern",
            "esrs_bezug": t["std"], "has_iro": "[" + ", ".join(t["iros"]) + "]"},
            f"# Thema: {t['sub']}\n\n**ESRS-Standard:** {t['std']} · "
            f"**Wesentlich:** {'Ja' if t['mat'] else 'Nein'} · **IROs:** {len(t['iros'])}\n\n"
            + ("Wesentlich laut DMA 2026." if t["mat"]
               else "Bewertet, aber **nicht wesentlich** — als Ausschluss-Begründung erfasst (Assurance)."))

    write("decision-dma-2026", {
        "id": "decision-dma-2026", "type": "decision", "owner": "person-dma-lead",
        "status": "beschlossen", "stand": STAND, "review_zyklus": "P12M", "vertraulichkeit": "intern",
        "esrs_bezug": "ESRS-2-IRO-1", "datum": "2026-06-09",
        "entscheidung": q(f"DMA 2026: {n_mat} von {len(rows)} IROs wesentlich, in {n_mat_topics} "
                          f"von {len(topics)} Themen."),
        "decided_by": "[person-dma-lead]", "based_on": "[methodology-dma-2026]",
        "affects": "[" + ", ".join(topics) + "]"},
        f"# Entscheidung: DMA-2026 Ergebnis\n\n{n_mat} von {len(rows)} IROs wesentlich, gruppiert in "
        f"{len(topics)} Themen ({n_mat_topics} wesentlich, {len(topics)-n_mat_topics} nicht wesentlich) "
        f"über die ESRS-Standards E1-E5, ES, G1, S1-S4.")

    # Sheet 2: Abhängigkeiten von natürlichem/sozialem Kapital je Standard
    ndep, seen = 0, {}
    for r in wb[wb.sheetnames[1]].iter_rows(values_only=True):
        c0 = str(r[0]).strip() if r and r[0] else ""
        if not re.fullmatch(r"(ES|E\d|S\d|G\d)", c0) or len(r) < 2 or not r[1]:
            continue
        ressource = str(r[1]).strip()
        seen[c0] = seen.get(c0, 0) + 1
        did = f"dependency-{slug(c0)}-{seen[c0]:02d}-{slug(ressource, 26)}"
        write(did, {
            "id": did, "type": "dependency", "owner": "person-dma-lead", "status": "erfasst",
            "stand": STAND, "vertraulichkeit": "intern", "esrs_bezug": c0,
            "ressource": q(ressource), "relevanz": q(r[3] if len(r) > 3 else "")},
            f"# Abhängigkeit: {ressource}\n\n**Standard:** {c0} · "
            f"**Relevanz (R/O):** {r[3] if len(r) > 3 and r[3] else '—'}\n\n"
            f"{r[2] if len(r) > 2 and r[2] else '—'}")
        ndep += 1

    print(f"Importiert: {len(topics)} Themen ({n_mat_topics} wesentlich) + {len(rows)} IROs "
          f"({n_mat} wesentlich) + {ndep} Dependencies + 3 Stamm "
          f"= {len(topics)+len(rows)+ndep+3} Objekte")


if __name__ == "__main__":
    main()
